"""
Offline unit tests for Module 2 (Source Summarizer).

Run from the backend/ folder:  python -m unittest discover -s tests -v
No network, API keys or real PDFs are needed.
"""

from __future__ import annotations

import asyncio
import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK

from app.config import Settings, get_settings
from app.main import app
from app.modules.research.services import library
from app.modules.summarize.schemas import SourceMeta
from app.modules.summarize.services.cite import build_citations
from app.modules.summarize.services.export import summary_docx
from app.modules.summarize.services.extract import (
    DocxError,
    PdfError,
    document_from_pages,
    document_from_text,
    read_docx,
    read_pdf,
    split_sentences,
)
from app.modules.summarize.services.save import save_summary
from app.modules.summarize.services.summarize import (
    _draft_from_llm,
    _extract_json,
    format_notes,
    heuristic_draft,
    summarize_document,
)

OFFLINE_SETTINGS = Settings(summarize_with_llm=False, llm_keyless_fallback=False)
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

ABSTRACT = """Background: Sleep quality is affected by ambient temperature, but the mechanism is unclear. We conducted a survey of 412 university students during a heatwave. Participants reported sleep duration and room temperature each night for two weeks. Results showed that each 1 degree C increase in bedroom temperature was associated with 7.4 minutes less sleep (p < 0.01). Students without air conditioning were significantly more likely to report poor sleep quality. However, the sample was limited to one university and relied on self-reported data. Future research should use wearable devices. Conclusions: Warmer bedrooms reduce sleep duration, and access to cooling is an important health equity issue."""


def build_summary() -> dict:
    document = document_from_text(ABSTRACT)
    draft = heuristic_draft(document, quote_count=3)
    return {
        "source": {
            "title": "Bedroom temperature and student sleep",
            "authors": ["Ada Lovelace", "Grace Hopper"],
            "year": 2025,
            "journal": "Journal of Sleep Studies",
            "doi": "10.1000/sleep.2025.1",
            "url": None,
        },
        "tldr": draft.tldr,
        "key_points": draft.key_points,
        "methods": draft.methods,
        "findings": draft.findings,
        "limitations": draft.limitations,
        "quotes": [
            {
                "text": document.sentences[index].text,
                "page": document.sentences[index].page,
                "note": "",
            }
            for index in draft.quote_indices
        ],
        "excerpt": document.excerpt,
        "ai_source": "heuristic",
    }


def build_docx() -> bytes:
    """A small Word file with a page break and a table."""
    document = DocxDocument()
    document.add_paragraph("We surveyed 412 students during a heatwave.")
    document.add_paragraph("Participants recorded sleep and bedroom temperature.")
    page_break = document.add_paragraph("Methods are described below.")
    page_break.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("Results showed a 7.4 minute reduction in sleep per degree.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Sample"
    table.cell(0, 1).text = "412 students"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class ExtractTests(unittest.TestCase):
    def test_split_keeps_abbreviations_together(self):
        sentences = split_sentences(
            "Smith et al. (2020) studied sleep. They found a small effect. "
            "The result was significant, e.g. p < 0.05."
        )
        self.assertEqual(len(sentences), 3, sentences)
        self.assertTrue(sentences[0].startswith("Smith et al. (2020)"))

    def test_page_numbers_survive(self):
        document = document_from_pages(
            [
                "We used a survey of 412 students. Results showed less sleep.",
                "However, the sample was limited to one university.",
            ]
        )
        self.assertEqual(document.page_count, 2)
        self.assertEqual([sentence.page for sentence in document.sentences], [1, 1, 2])

    def test_pdf_magic_and_blank_pdf_are_rejected(self):
        with self.assertRaises(PdfError):
            read_pdf(b"not really a pdf")

        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        buffer = io.BytesIO()
        writer.write(buffer)
        with self.assertRaises(PdfError) as caught:
            read_pdf(buffer.getvalue())
        self.assertIn("selectable text", str(caught.exception))

    def test_docx_extracts_text_tables_and_page_breaks(self):
        document = read_docx(build_docx())
        self.assertEqual(document.kind, "docx")
        self.assertEqual(document.page_count, 2)
        self.assertIn("412 students", document.text)
        self.assertIn("Sample | 412 students", document.text)
        self.assertIn("7.4 minute", document.text)
        self.assertEqual({sentence.page for sentence in document.sentences}, {1, 2})
        page_one = [s.text for s in document.sentences if s.page == 1]
        self.assertTrue(any("surveyed 412" in sentence for sentence in page_one))

    def test_docx_without_page_breaks_keeps_honest_page_numbers(self):
        document = DocxDocument()
        document.add_paragraph(
            "We surveyed 412 students during a heatwave and measured their sleep."
        )
        buffer = io.BytesIO()
        document.save(buffer)

        parsed = read_docx(buffer.getvalue())
        self.assertIsNone(parsed.page_count)
        self.assertTrue(all(sentence.page is None for sentence in parsed.sentences))
        self.assertTrue(any("page numbers" in warning for warning in parsed.warnings))

    def test_docx_rejects_empty_and_wrong_files(self):
        with self.assertRaises(DocxError):
            read_docx(b"")
        with self.assertRaises(DocxError):
            read_docx(b"not a docx")
        with self.assertRaises(DocxError):
            read_docx(b"\xd0\xcf\x11\xe0old-word-file")

        empty = DocxDocument()
        buffer = io.BytesIO()
        empty.save(buffer)
        with self.assertRaises(DocxError) as caught:
            read_docx(buffer.getvalue())
        self.assertIn("no selectable text", str(caught.exception))


class SummaryTests(unittest.TestCase):
    def test_heuristic_has_every_section(self):
        document = document_from_text(ABSTRACT)
        draft = heuristic_draft(document, quote_count=3)

        self.assertTrue(draft.tldr)
        self.assertTrue(draft.key_points)
        self.assertTrue(draft.methods)
        self.assertTrue(draft.findings)
        self.assertTrue(draft.limitations)
        self.assertEqual(len(draft.quote_indices), 3)
        self.assertIn("sample was limited", " ".join(draft.limitations))

    def test_quotes_are_verbatim_and_keep_pages(self):
        document = document_from_pages(
            [
                "We conducted a survey of 412 students during a heatwave. "
                "The method section explains the sampling.",
                "Results showed a 7.4 minute reduction in sleep. "
                "However, the sample was limited to one university.",
            ]
        )
        draft = heuristic_draft(document, quote_count=2)
        self.assertEqual(len(draft.quote_indices), 2)
        for index in draft.quote_indices:
            sentence = next(item for item in document.sentences if item.index == index)
            self.assertIn(sentence.text, [item.text for item in document.sentences])
            self.assertIn(sentence.page, (1, 2))

    def test_offline_summary_never_uses_network(self):
        document = document_from_text(ABSTRACT)
        draft, kind = asyncio.run(
            summarize_document(
                document,
                source_title="Sleep",
                settings=OFFLINE_SETTINGS,
                quote_count=2,
            )
        )
        self.assertEqual(kind, "heuristic")
        self.assertEqual(draft.ai_source, "heuristic")

    def test_llm_quote_numbers_map_to_real_sentences(self):
        document = document_from_text(ABSTRACT)
        answer = """```json
        {
          "tldr": "Warmer bedrooms shortened sleep in a student survey.",
          "key_points": ["Heat harms sleep."],
          "methods": ["Survey of 412 students."],
          "findings": ["Sleep fell by 7.4 minutes per degree."],
          "limitations": ["One university, self-report."],
          "quotes": [4, 99, 4]
        }
        ```"""
        draft = _draft_from_llm(answer, document, quote_count=3)
        self.assertIsNotNone(draft)
        assert draft is not None
        self.assertEqual(draft.ai_source, "llm")
        self.assertEqual(draft.quote_indices, [3])
        self.assertIn("7.4 minutes", document.sentences[3].text)

    def test_detail_level_changes_how_much_the_llm_returns(self):
        document = document_from_text(ABSTRACT)
        answer = json.dumps(
            {
                "tldr": "Heat reduces sleep.",
                "key_points": [f"Point {index}" for index in range(1, 7)],
                "methods": [f"Method {index}" for index in range(1, 6)],
                "findings": [f"Finding {index}" for index in range(1, 6)],
                "limitations": [f"Limit {index}" for index in range(1, 6)],
                "quotes": [],
            }
        )
        quick = _draft_from_llm(answer, document, quote_count=0, detail="quick")
        standard = _draft_from_llm(answer, document, quote_count=0)
        deep = _draft_from_llm(answer, document, quote_count=0, detail="deep")
        assert quick is not None and standard is not None and deep is not None
        self.assertEqual(len(quick.key_points), 3)
        self.assertEqual(len(standard.key_points), 5)
        self.assertEqual(len(deep.key_points), 6)
        self.assertEqual(len(quick.methods), 2)
        self.assertEqual(len(deep.methods), 4)

    def test_heuristic_detail_never_exceeds_its_limits(self):
        document = document_from_text(ABSTRACT)
        quick = heuristic_draft(document, quote_count=2, detail="quick")
        deep = heuristic_draft(document, quote_count=2, detail="deep")
        self.assertLessEqual(len(quick.key_points), 3)
        self.assertLessEqual(len(quick.methods), 2)
        self.assertLessEqual(len(quick.findings), 2)
        self.assertLessEqual(len(quick.limitations), 2)
        self.assertLessEqual(len(deep.key_points), 5)
        self.assertLessEqual(len(deep.findings), 4)

    def test_extract_json_rejects_prose(self):
        self.assertIsNone(_extract_json("I could not summarize this."))
        self.assertEqual(_extract_json('{"tldr": "ok"}'), {"tldr": "ok"})


class _FakeLLM:
    """Stands in for LLMClient so repair/fallback paths stay offline."""

    def __init__(self, replies: list[str]) -> None:
        self.replies = list(replies)
        self.calls = 0

    def __call__(self, settings: Settings) -> "_FakeLLM":
        return self

    @property
    def available(self) -> bool:
        return True

    async def complete(self, **kwargs: object) -> str:
        self.calls += 1
        if not self.replies:
            raise AssertionError("the AI was called more times than the test allowed")
        return self.replies.pop(0)


class RepairTests(unittest.TestCase):
    def test_malformed_first_reply_is_retried_once(self):
        document = document_from_text(ABSTRACT)
        valid = json.dumps(
            {
                "tldr": "Warmer bedrooms shortened sleep.",
                "key_points": ["Heat reduces sleep."],
                "methods": ["Survey of 412 students."],
                "findings": ["Sleep fell by 7.4 minutes per degree."],
                "limitations": ["One university."],
                "quotes": [],
            }
        )
        fake = _FakeLLM(["not json at all", valid])
        settings = Settings(summarize_with_llm=True, llm_keyless_fallback=False)
        with mock.patch(
            "app.modules.summarize.services.summarize.LLMClient", fake
        ):
            draft, kind = asyncio.run(
                summarize_document(
                    document, source_title="Sleep", settings=settings, quote_count=0
                )
            )
        self.assertEqual(kind, "llm")
        self.assertEqual(fake.calls, 2)
        self.assertTrue(any("first reply" in warning for warning in draft.warnings))

    def test_two_bad_replies_fall_back_with_a_visible_warning(self):
        document = document_from_text(ABSTRACT)
        fake = _FakeLLM(["still not json", "also not json"])
        settings = Settings(summarize_with_llm=True, llm_keyless_fallback=True)
        with mock.patch(
            "app.modules.summarize.services.summarize.LLMClient", fake
        ):
            draft, kind = asyncio.run(
                summarize_document(
                    document, source_title="Sleep", settings=settings, quote_count=2
                )
            )
        self.assertEqual(kind, "heuristic")
        self.assertEqual(fake.calls, 2)
        self.assertTrue(any("could not be read" in warning for warning in draft.warnings))


class CitationTests(unittest.TestCase):
    def test_two_authors_get_apa_in_text_and_reference(self):
        citations = build_citations(
            SourceMeta(
                title="Bedroom temperature and student sleep",
                authors=["Ada Lovelace", "Grace Hopper"],
                year=2025,
                journal="Journal of Sleep Studies",
                doi="https://doi.org/10.1000/sleep.2025.1",
            )
        )
        self.assertEqual(citations.author_label, "Lovelace & Hopper")
        self.assertEqual(citations.in_text, "(Lovelace & Hopper, 2025)")
        self.assertEqual(citations.narrative, "Lovelace and Hopper (2025)")
        self.assertIn("Lovelace, A., & Hopper, G. (2025).", citations.reference)
        self.assertIn("https://doi.org/10.1000/sleep.2025.1", citations.reference)
        self.assertNotIn("https://doi.org/https://doi.org", citations.reference)
        self.assertIn("Draft built", citations.note)

    def test_three_authors_use_et_al_and_missing_year_uses_nd(self):
        citations = build_citations(
            SourceMeta(
                title="A study",
                authors=["Ana One", "Ben Two", "Cleo Three"],
                year=None,
            )
        )
        self.assertEqual(citations.author_label, "One et al.")
        self.assertEqual(citations.in_text, "(One et al., n.d.)")

    def test_missing_authors_does_not_invent_a_reference(self):
        citations = build_citations(SourceMeta(title="A paper", year=2025))
        self.assertEqual(citations.in_text, "")
        self.assertEqual(citations.reference, "")
        self.assertIn("at least one author", citations.note)

    def test_missing_title_keeps_in_text_but_skips_reference(self):
        citations = build_citations(SourceMeta(authors=["Ada Lovelace"], year=2025))
        self.assertEqual(citations.in_text, "(Lovelace, 2025)")
        self.assertEqual(citations.reference, "")
        self.assertIn("title", citations.note)


class ExportTests(unittest.TestCase):
    def test_summary_docx_contains_the_study_sections(self):
        summary = build_summary()
        summary["citations"] = {
            "reference": "Lovelace, A. (2025). Bedroom temperature. https://doi.org/10.1000/x"
        }
        data = summary_docx(summary)
        self.assertTrue(data.startswith(b"PK"))

        document = DocxDocument(io.BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Bedroom temperature and student sleep", text)
        self.assertIn("TL;DR", text)
        self.assertIn("Key points", text)
        self.assertIn("Verbatim quotes", text)
        self.assertIn("APA reference", text)
        self.assertIn("Generated by Scholaris", text)


class SaveTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db = Path(self.temp.name) / "library.sqlite3"
        library.init_db(self.db)

    def tearDown(self):
        self.temp.cleanup()

    def test_save_creates_one_library_entry_and_is_idempotent(self):
        summary = build_summary()
        first = save_summary(self.db, summary)
        second = save_summary(self.db, summary)

        self.assertTrue(first["created"])
        self.assertFalse(second["created"])
        self.assertEqual(first["library_id"], second["library_id"])

        items = library.list_sources(self.db)
        self.assertEqual(len(items), 1)
        item = items[0]
        self.assertIn("TL;DR", item["notes"])
        self.assertEqual(item["notes"].count("SUMMARY —"), 1)
        self.assertIn("summary", item["tags"])
        self.assertEqual(item["citation_styles"]["apa"].count("Lovelace"), 1)

    def test_save_appends_to_an_existing_saved_source(self):
        summary = build_summary()
        library.save_source(
            self.db,
            {
                "id": "doi:10.1000/sleep.2025.1",
                "title": summary["source"]["title"],
                "authors": summary["source"]["authors"],
                "year": 2025,
                "journal": summary["source"]["journal"],
                "doi": summary["source"]["doi"],
                "citation_styles": {"apa": "", "mla": "", "chicago": ""},
                "credibility": [],
            },
            notes="My own note.",
            tags=["key-source"],
        )
        result = save_summary(self.db, summary, library_id="doi:10.1000/sleep.2025.1")
        self.assertFalse(result["created"])
        item = library.get_source(self.db, "doi:10.1000/sleep.2025.1")
        assert item is not None
        self.assertTrue(item["notes"].startswith("My own note."))
        self.assertIn("TL;DR", item["notes"])
        self.assertEqual(set(item["tags"]), {"key-source", "summary"})

    def test_missing_library_id_raises(self):
        with self.assertRaises(KeyError):
            save_summary(self.db, build_summary(), library_id="doi:missing")

    def test_format_notes_includes_page_numbers(self):
        summary = build_summary()
        summary["quotes"] = [{"text": "A verbatim quote.", "page": 4, "note": ""}]
        notes = format_notes(summary)
        self.assertIn('"A verbatim quote." (p. 4)', notes)

    def test_stashed_text_is_saved_once_and_kept_on_resave(self):
        summary = build_summary()
        summary["summary_id"] = library.stash_source_text(
            self.db, "The full source text lives here."
        )
        result = save_summary(self.db, summary)
        self.assertEqual(
            library.get_source_text(self.db, result["library_id"]),
            "The full source text lives here.",
        )
        # The token is single-use; a second save must not wipe the stored text.
        self.assertEqual(library.take_source_text(self.db, summary["summary_id"]), "")
        save_summary(self.db, summary)
        self.assertEqual(
            library.get_source_text(self.db, result["library_id"]),
            "The full source text lives here.",
        )

    def test_init_db_migrates_a_v1_database(self):
        path = Path(self.temp.name) / "old.sqlite3"
        connection = sqlite3.connect(path)
        connection.execute(
            """
            CREATE TABLE sources (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                year INTEGER,
                authors TEXT NOT NULL DEFAULT '[]',
                journal TEXT,
                doi TEXT,
                url TEXT,
                payload TEXT NOT NULL,
                notes TEXT NOT NULL DEFAULT '',
                tags TEXT NOT NULL DEFAULT '[]',
                saved_at TEXT NOT NULL
            )
            """
        )
        connection.execute(
            "INSERT INTO sources (id, title, authors, payload, saved_at) "
            "VALUES ('doi:x', 'Old source', '[]', '{}', '2026-01-01T00:00:00+00:00')"
        )
        connection.commit()
        connection.close()

        library.init_db(path)
        self.assertEqual(library.get_source_text(path, "doi:x"), "")
        self.assertTrue(library.set_source_text(path, "doi:x", "Recovered text."))
        self.assertEqual(library.get_source_text(path, "doi:x"), "Recovered text.")

    def test_unknown_stash_token_is_safe(self):
        self.assertEqual(library.take_source_text(self.db, ""), "")
        self.assertEqual(library.take_source_text(self.db, "does-not-exist"), "")


class RouteTests(unittest.TestCase):
    def setUp(self):
        from fastapi.testclient import TestClient

        self.temp = tempfile.TemporaryDirectory()
        self.db = Path(self.temp.name) / "library.sqlite3"
        library.init_db(self.db)
        self.settings = Settings(
            database_path=self.db,
            summarize_with_llm=False,
            llm_keyless_fallback=False,
        )
        app.dependency_overrides[get_settings] = lambda: self.settings
        self.client = TestClient(app)

    def tearDown(self):
        app.dependency_overrides.clear()
        self.client.close()
        self.temp.cleanup()

    def _summarize(self) -> dict:
        response = self.client.post(
            "/api/summarize/text",
            json={
                "text": ABSTRACT,
                "quote_count": 2,
                "source": {
                    "title": "Bedroom temperature and student sleep",
                    "authors": ["Ada Lovelace"],
                    "year": 2025,
                    "journal": "Journal of Sleep Studies",
                    "doi": "10.1000/sleep.2025.1",
                    "url": None,
                },
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_text_summary_endpoint(self):
        body = self._summarize()
        self.assertEqual(body["ai_source"], "heuristic")
        self.assertTrue(body["tldr"])
        self.assertTrue(body["findings"])
        self.assertEqual(len(body["quotes"]), 2)
        self.assertIn("page", body["quotes"][0])

    def test_short_text_is_rejected(self):
        response = self.client.post(
            "/api/summarize/text",
            json={"text": "Too short to summarize.", "source": {}, "quote_count": 2},
        )
        self.assertEqual(response.status_code, 422)

    def test_save_endpoint_writes_into_library(self):
        summary = self._summarize()
        response = self.client.post("/api/summarize/save", json={"summary": summary})
        self.assertEqual(response.status_code, 200, response.text)
        saved = response.json()
        self.assertTrue(saved["created"])

        library_response = self.client.get("/api/research/library")
        self.assertEqual(library_response.status_code, 200)
        items = library_response.json()["sources"]
        self.assertEqual(len(items), 1)
        self.assertIn("TL;DR", items[0]["notes"])
        self.assertEqual(items[0]["library_id"], saved["library_id"])

    def test_pdf_upload_rejects_non_pdf(self):
        response = self.client.post(
            "/api/summarize/pdf",
            files={"file": ("paper.pdf", b"hello", "application/pdf")},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("not a PDF", response.json()["detail"])

    def test_docx_upload_honours_detail_and_plain_mode(self):
        response = self.client.post(
            "/api/summarize/docx",
            files={"file": ("paper.docx", build_docx(), DOCX_MIME)},
            data={
                "detail": "deep",
                "mode": "plain",
                "quote_count": "2",
                "title": "Sleep and heat",
                "authors": "Ada Lovelace",
                "year": "2025",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["detail"], "deep")
        self.assertEqual(body["mode"], "plain")
        self.assertEqual(body["page_count"], 2)
        self.assertEqual(body["citations"]["author_label"], "Lovelace")
        self.assertEqual(body["citations"]["in_text"], "(Lovelace, 2025)")
        self.assertTrue(body["summary_id"])
        self.assertEqual(len(body["quotes"]), 2)
        self.assertTrue(any(quote["page"] in (1, 2) for quote in body["quotes"]))

    def test_docx_upload_rejects_a_non_docx_file(self):
        response = self.client.post(
            "/api/summarize/docx",
            files={"file": ("paper.docx", b"hello", DOCX_MIME)},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("not a .docx", response.json()["detail"])

    def test_doi_lookup_rejects_an_invalid_doi_without_network(self):
        response = self.client.get("/api/summarize/doi", params={"doi": "not-a-doi"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("DOI", response.json()["detail"])

    def test_export_docx_returns_a_word_document(self):
        summary = self._summarize()
        response = self.client.post("/api/summarize/export/docx", json=summary)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn("wordprocessingml", response.headers["content-type"])
        self.assertIn("attachment;", response.headers["content-disposition"])

        document = DocxDocument(io.BytesIO(response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Bedroom temperature and student sleep", text)
        self.assertIn("APA reference", text)
        self.assertIn("Generated by Scholaris", text)

    def test_save_endpoint_keeps_the_source_text_private(self):
        summary = self._summarize()
        self.assertTrue(summary["summary_id"])
        response = self.client.post("/api/summarize/save", json={"summary": summary})
        self.assertEqual(response.status_code, 200, response.text)
        saved = response.json()
        self.assertEqual(library.get_source_text(self.db, saved["library_id"]), ABSTRACT)
        self.assertEqual(library.take_source_text(self.db, summary["summary_id"]), "")

        item = self.client.get("/api/research/library").json()["sources"][0]
        self.assertNotIn("full_text", item)

    def test_text_summary_accepts_detail_and_mode(self):
        response = self.client.post(
            "/api/summarize/text",
            json={
                "text": ABSTRACT,
                "quote_count": 1,
                "detail": "quick",
                "mode": "plain",
                "source": {"title": "Sleep and heat"},
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["detail"], "quick")
        self.assertEqual(body["mode"], "plain")
        self.assertEqual(body["citations"]["in_text"], "")
        self.assertIn("author", body["citations"]["note"])
        self.assertLessEqual(len(body["key_points"]), 3)


if __name__ == "__main__":
    unittest.main()
