"""
Personal source library + bibliography export.

SQLite keeps this dependency-free and portable. Each saved source stores the
full API payload, so the library still renders correctly even if an external
API changes its response shape later.
"""

from __future__ import annotations

import io
import json
import csv
import re
import sqlite3
import unicodedata
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator

from .citations import split_name
from .text import clean_text, expand_page_range, normalize_journal_name, split_volume_issue

SCHEMA = """
CREATE TABLE IF NOT EXISTS sources (
    id        TEXT PRIMARY KEY,
    title     TEXT NOT NULL,
    year      INTEGER,
    authors   TEXT NOT NULL DEFAULT '[]',
    journal   TEXT,
    doi       TEXT,
    url       TEXT,
    payload   TEXT NOT NULL,
    notes     TEXT NOT NULL DEFAULT '',
    tags      TEXT NOT NULL DEFAULT '[]',
    saved_at  TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_sources_saved_at ON sources(saved_at DESC);
"""


def init_db(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with _connect(path) as connection:
        connection.executescript(SCHEMA)


@contextmanager
def _connect(path: Path) -> Iterator[sqlite3.Connection]:
    connection = sqlite3.connect(path, timeout=10)
    connection.row_factory = sqlite3.Row
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _row_to_item(row: sqlite3.Row) -> dict[str, Any]:
    payload = json.loads(row["payload"] or "{}")
    # Older saves may hold provider quirks (journal abbreviations, abbreviated
    # page ranges, a stray space before a colon). Normalise on read so the
    # library and its exports heal without a database migration.
    payload["title"] = clean_text(payload.get("title")) or payload.get("title") or ""
    payload["journal"] = normalize_journal_name(payload.get("journal"))
    restored_pages = expand_page_range(payload.get("pages"))
    if restored_pages:
        payload["pages"] = restored_pages
    volume = clean_text(payload.get("volume")) or None
    issue = clean_text(payload.get("issue")) or None
    if volume and not issue:
        volume, issue = split_volume_issue(volume)
    payload["volume"], payload["issue"] = volume, issue
    payload.update(
        {
            "library_id": row["id"],
            "notes": row["notes"] or "",
            "tags": json.loads(row["tags"] or "[]"),
            "saved_at": row["saved_at"],
        }
    )
    return payload


def save_source(path: Path, source: dict[str, Any], notes: str = "", tags: list[str] | None = None) -> dict[str, Any]:
    library_id = str(source["id"])
    tags = tags or []
    with _connect(path) as connection:
        connection.execute(
            """
            INSERT INTO sources (id, title, year, authors, journal, doi, url, payload, notes, tags, saved_at)
            VALUES (:id, :title, :year, :authors, :journal, :doi, :url, :payload, :notes, :tags, :saved_at)
            ON CONFLICT(id) DO UPDATE SET
                title   = excluded.title,
                year    = excluded.year,
                authors = excluded.authors,
                journal = excluded.journal,
                doi     = excluded.doi,
                url     = excluded.url,
                payload = excluded.payload,
                notes   = CASE WHEN excluded.notes <> '' THEN excluded.notes ELSE sources.notes END,
                tags    = CASE WHEN excluded.tags <> '[]' THEN excluded.tags ELSE sources.tags END
            """,
            {
                "id": library_id,
                "title": source.get("title") or "Untitled",
                "year": source.get("year"),
                "authors": json.dumps(source.get("authors") or []),
                "journal": source.get("journal"),
                "doi": source.get("doi"),
                "url": source.get("url"),
                "payload": json.dumps(source),
                "notes": notes or "",
                "tags": json.dumps(tags),
                "saved_at": _now(),
            },
        )
    item = next((entry for entry in list_sources(path) if entry["library_id"] == library_id), None)
    return item or {}


def list_sources(path: Path) -> list[dict[str, Any]]:
    with _connect(path) as connection:
        rows = connection.execute("SELECT * FROM sources ORDER BY saved_at DESC").fetchall()
    return [_row_to_item(row) for row in rows]


def get_source(path: Path, library_id: str) -> dict[str, Any] | None:
    with _connect(path) as connection:
        row = connection.execute("SELECT * FROM sources WHERE id = ?", (library_id,)).fetchone()
    return _row_to_item(row) if row else None


def update_source(path: Path, library_id: str, *, notes: str | None, tags: list[str] | None) -> dict[str, Any] | None:
    with _connect(path) as connection:
        row = connection.execute("SELECT * FROM sources WHERE id = ?", (library_id,)).fetchone()
        if row is None:
            return None
        new_notes = row["notes"] if notes is None else notes
        new_tags = row["tags"] if tags is None else json.dumps(tags)
        connection.execute(
            "UPDATE sources SET notes = ?, tags = ? WHERE id = ?",
            (new_notes, new_tags, library_id),
        )
        updated = connection.execute("SELECT * FROM sources WHERE id = ?", (library_id,)).fetchone()
    return _row_to_item(updated) if updated else None


def delete_source(path: Path, library_id: str) -> bool:
    with _connect(path) as connection:
        cursor = connection.execute("DELETE FROM sources WHERE id = ?", (library_id,))
    return cursor.rowcount > 0


# --------------------------------------------------------------------- exports


def _slug(value: str) -> str:
    text = unicodedata.normalize("NFKD", value)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-zA-Z0-9]+", "", text)
    return text.lower()


def _bib_escape(value: Any) -> str:
    text = str(value or "")
    for character in "&%$#_":
        text = text.replace(character, f"\\{character}")
    return text.replace("{", "\\{").replace("}", "\\}")


def _bib_authors(item: dict[str, Any]) -> str:
    structured = item.get("author_details") or []
    formatted: list[str] = []
    if structured:
        for author in structured:
            if author.get("literal"):
                formatted.append(author["literal"])
                continue
            family = author.get("family") or ""
            given = author.get("given") or ""
            suffix = author.get("suffix") or ""
            parts = [part for part in (family, suffix, given) if part]
            if parts:
                formatted.append(", ".join(parts))
        return " and ".join(formatted)

    for name in item.get("authors") or []:
        family, given = split_name(name)
        formatted.append(f"{family}, {given}".strip(", "))
    return " and ".join(formatted)


def export_bibtex(items: list[dict[str, Any]]) -> str:
    used_keys: set[str] = set()
    entries: list[str] = []

    for item in items:
        authors = item.get("authors") or []
        structured = item.get("author_details") or []
        if structured:
            first = structured[0]
            family = _slug(first.get("family") or first.get("literal") or "") or "anon"
        else:
            family = _slug(split_name(authors[0])[0]) if authors else "anon"
        year = item.get("year") or "nd"
        first_word = _slug((item.get("title") or "").split(" ")[0]) or "work"
        base_key = f"{family}{year}{first_word}" or "source"
        key = base_key
        suffix = 1
        while key in used_keys:
            suffix += 1
            key = f"{base_key}{chr(96 + suffix)}"
        used_keys.add(key)

        fields: list[tuple[str, str]] = []
        if item.get("title"):
            fields.append(("title", "{" + _bib_escape(item["title"]) + "}"))
        bib_authors = _bib_authors(item)
        if bib_authors:
            fields.append(("author", _bib_escape(bib_authors)))
        if item.get("year"):
            fields.append(("year", str(item["year"])))
        if item.get("journal"):
            fields.append(("journal", _bib_escape(item["journal"])))
        if item.get("volume"):
            fields.append(("volume", _bib_escape(item["volume"])))
        if item.get("issue"):
            fields.append(("number", _bib_escape(item["issue"])))
        if item.get("pages"):
            fields.append(("pages", _bib_escape(item["pages"])))
        elif item.get("article_number"):
            fields.append(("eid", _bib_escape(item["article_number"])))
        if item.get("publisher"):
            fields.append(("publisher", _bib_escape(item["publisher"])))
        if item.get("doi"):
            fields.append(("doi", _bib_escape(item["doi"])))
        if item.get("url"):
            fields.append(("url", _bib_escape(item["url"])))
        author_check = item.get("author_check") or {}
        if author_check.get("status") == "needs_checking":
            reason = author_check.get("reason") or "Provider records disagree."
            fields.append(
                ("note", _bib_escape(f"NEEDS CHECKING: author names. {reason}"))
            )
        tags = item.get("tags") or []
        if tags:
            fields.append(("keywords", _bib_escape(", ".join(tags))))

        entry_type = "article" if item.get("journal") else "misc"
        body = ",\n  ".join(f"{name} = {{{value}}}" for name, value in fields)
        entries.append(f"@{entry_type}{{{key},\n  {body}\n}}")

    return "\n\n".join(entries) + ("\n" if entries else "")


def export_txt(items: list[dict[str, Any]]) -> str:
    lines = [
        "Scholaris — Bibliography",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"Sources: {len(items)}",
        "",
    ]
    for index, item in enumerate(items, start=1):
        apa = (item.get("citation_styles") or {}).get("apa") or item.get("title", "")
        lines.append(f"{index}. {apa}")
        if item.get("notes"):
            lines.append(f"   Notes: {item['notes']}")
        tags = item.get("tags") or []
        if tags:
            lines.append(f"   Tags: {', '.join(tags)}")
        lines.append("")
    return "\n".join(lines)


def export_csv(items: list[dict[str, Any]]) -> str:
    """A spreadsheet-friendly export of the current result set or library."""
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "Title",
            "Authors",
            "Year",
            "Journal",
            "DOI",
            "URL",
            "Citations",
            "APA",
            "Notes",
            "Tags",
        ]
    )
    for item in items:
        apa = (item.get("citation_styles") or {}).get("apa") or ""
        writer.writerow(
            [
                item.get("title") or "",
                "; ".join(item.get("authors") or []),
                item.get("year") or "",
                item.get("journal") or "",
                item.get("doi") or "",
                item.get("url") or "",
                item.get("citation_count") or "",
                apa,
                item.get("notes") or "",
                ", ".join(item.get("tags") or []),
            ]
        )
    return buffer.getvalue()


def export_docx(items: list[dict[str, Any]]) -> bytes:
    # Imported here so the rest of the app works even without python-docx.
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    document.add_heading("Bibliography", level=1)
    document.add_paragraph(
        f"Generated by Scholaris on {datetime.now().strftime('%Y-%m-%d')} · {len(items)} sources"
    )

    for item in items:
        apa = (item.get("citation_styles") or {}).get("apa") or item.get("title", "")
        paragraph = document.add_paragraph(apa)
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        if item.get("notes"):
            note = document.add_paragraph(f"Notes: {item['notes']}")
            note.runs[0].italic = True
            note.runs[0].font.size = Pt(9)
        tags = item.get("tags") or []
        if tags:
            tag_paragraph = document.add_paragraph(f"Tags: {', '.join(tags)}")
            tag_paragraph.runs[0].italic = True
            tag_paragraph.runs[0].font.size = Pt(9)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
