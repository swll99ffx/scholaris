"""
Turn pasted text or a PDF into clean sentences that keep their page numbers.

The rest of Module 2 only works with `Document`, so the summary generator and
the quote finder never need to know whether the words came from a textarea or
from page 4 of a PDF.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

# Guard rails so a huge file cannot make a request slow or fill the RAM.
MAX_PDF_PAGES = 60
MAX_PDF_BYTES = 15 * 1024 * 1024
MAX_DOCX_BYTES = 15 * 1024 * 1024
MAX_DOCX_XML_BYTES = 40 * 1024 * 1024
MAX_DOCX_TOTAL_BYTES = 120 * 1024 * 1024
MAX_TEXT_CHARS = 120_000
MIN_PDF_TEXT_CHARS = 40
MIN_DOCX_TEXT_CHARS = 40

_WHITESPACE_RE = re.compile(r"[ \t\f\v\u00a0]+")
_SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\"'“”‘’(])")
_INLINE_CITATION_RE = re.compile(r"\s*\[(?:\d+[\s,;–-]*)+]")
_LIGATURES = {
    "\ufb00": "ff",
    "\ufb01": "fi",
    "\ufb02": "fl",
    "\ufb03": "ffi",
    "\ufb04": "ffl",
    "\ufb05": "st",
    "\ufb06": "st",
}
_ABBREVIATIONS = (
    "e.g.",
    "i.e.",
    "et al.",
    "vs.",
    "cf.",
    "approx.",
    "fig.",
    "figs.",
    "eq.",
    "no.",
    "dr.",
    "prof.",
    "st.",
    "inc.",
    "ltd.",
    "jr.",
    "sr.",
    "etc.",
    "vol.",
)


@dataclass(slots=True)
class Sentence:
    text: str
    page: int | None
    index: int


@dataclass(slots=True)
class Document:
    sentences: list[Sentence]
    text: str
    word_count: int
    kind: str = "text"
    page_count: int | None = None
    truncated: bool = False
    warnings: list[str] = field(default_factory=list)

    @property
    def excerpt(self) -> str:
        return excerpt(self.text)


class PdfError(ValueError):
    """A PDF could not be read (encrypted, scanned, not a PDF, ...)."""


class DocxError(ValueError):
    """A .docx could not be read (old .doc format, damaged, images only, ...)."""


def _normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    for ligature, replacement in _LIGATURES.items():
        text = text.replace(ligature, replacement)
    text = text.replace("\u00ad", "").replace("\r\n", "\n").replace("\r", "\n")
    # Re-join words that were hyphenated across a line break.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # A single newline is usually a wrapped line; keep blank lines as breaks.
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    text = re.sub(r"\n{2,}", "\n", text)
    text = _WHITESPACE_RE.sub(" ", text)
    return text.strip()


def clean_text(text: str) -> str:
    """Public helper for already-extracted text."""
    return _normalize(text)


def _is_abbreviation(sentence: str) -> bool:
    lowered = sentence.lower()
    return any(lowered.endswith(abbreviation) for abbreviation in _ABBREVIATIONS)


def split_sentences(text: str) -> list[str]:
    """Split clean prose into sentences without breaking common abbreviations."""
    text = clean_text(text)
    if not text:
        return []

    chunks = [chunk.strip() for chunk in _SENTENCE_SPLIT_RE.split(text) if chunk.strip()]
    sentences: list[str] = []
    for chunk in chunks:
        if sentences and (
            _is_abbreviation(sentences[-1])
            or len(chunk.split()) < 4
            or chunk[0].islower()
        ):
            sentences[-1] = f"{sentences[-1]} {chunk}"
        else:
            sentences.append(chunk)

    # A trailing fragment with no terminal punctuation is still useful.
    return [sentence for sentence in sentences if sentence]


def sentences_from_text(text: str, *, offset: int = 0) -> list[Sentence]:
    return [
        Sentence(text=sentence, page=None, index=offset + position)
        for position, sentence in enumerate(split_sentences(text))
    ]


def document_from_text(text: str) -> Document:
    cleaned = clean_text(text)
    if not cleaned:
        raise ValueError("The text was empty after cleaning. Paste a little more of the source.")
    truncated = len(cleaned) > MAX_TEXT_CHARS
    cleaned = cleaned[:MAX_TEXT_CHARS]
    sentences = sentences_from_text(cleaned)
    if not sentences:
        raise ValueError("No complete sentences were found in the text.")
    return Document(
        sentences=sentences,
        text=cleaned,
        word_count=len(cleaned.split()),
        kind="text",
        truncated=truncated,
        warnings=(
            ["Only the first part of the text was summarized (the limit is 120,000 characters)."]
            if truncated
            else []
        ),
    )


def document_from_pages(pages: list[str], *, kind: str = "pdf") -> Document:
    """Build a document from per-page strings, preserving page numbers."""
    sentences: list[Sentence] = []
    kept_pages: list[str] = []
    total_chars = 0
    truncated = False

    for page_number, raw_page in enumerate(pages, start=1):
        cleaned = clean_text(raw_page)
        if not cleaned:
            continue
        remaining = MAX_TEXT_CHARS - total_chars
        if remaining <= 0:
            truncated = True
            break
        if len(cleaned) > remaining:
            cleaned = cleaned[:remaining]
            truncated = True
        kept_pages.append(cleaned)
        total_chars += len(cleaned)
        offset = len(sentences)
        sentences.extend(
            Sentence(text=sentence, page=page_number, index=offset + position)
            for position, sentence in enumerate(split_sentences(cleaned))
        )

    if not sentences:
        raise PdfError(
            "This PDF has no selectable text. It may be a scan or an image export — "
            "paste the text instead, or run OCR on the file first."
        )

    warnings: list[str] = []
    if truncated:
        warnings.append(
            "Only part of the PDF was summarized (the limit is 120,000 characters)."
        )
    return Document(
        sentences=sentences,
        text=" ".join(kept_pages),
        word_count=sum(len(sentence.text.split()) for sentence in sentences),
        kind=kind,
        page_count=len(pages),
        truncated=truncated,
        warnings=warnings,
    )


def read_pdf(data: bytes) -> Document:
    """Extract text page by page from a PDF using pypdf (lazy import)."""
    if not data:
        raise PdfError("The uploaded file was empty.")
    if len(data) > MAX_PDF_BYTES:
        raise PdfError("That PDF is larger than 15 MB. Try a smaller file or paste the text.")
    if not data.lstrip().startswith(b"%PDF"):
        raise PdfError("That file is not a PDF. Upload a .pdf file or paste the text instead.")

    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is in requirements.txt
        raise PdfError("PDF support is not installed on the server (missing pypdf).") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise PdfError("That PDF is password-protected. Export an unlocked copy and try again.")
        pages = [page.extract_text() or "" for page in reader.pages[:MAX_PDF_PAGES]]
    except PdfError:
        raise
    except Exception as exc:  # noqa: BLE001 - pypdf raises several different errors
        raise PdfError("That PDF could not be read. It may be damaged or use an unsupported format.") from exc

    if len(reader.pages) > MAX_PDF_PAGES:
        document = document_from_pages(pages)
        document.page_count = len(reader.pages)
        document.warnings.append(
            f"Only the first {MAX_PDF_PAGES} pages were read (the PDF has {len(reader.pages)})."
        )
        document.truncated = True
        return document

    if sum(len(page.strip()) for page in pages) < MIN_PDF_TEXT_CHARS:
        raise PdfError(
            "This PDF has almost no selectable text. It may be a scan — paste the text instead, "
            "or run OCR on the file first."
        )
    return document_from_pages(pages)


def _paragraph_has_page_break(paragraph: object) -> bool:
    """True when a Word paragraph contains an explicit or rendered page break."""
    from docx.oxml.ns import qn

    element = paragraph._p  # type: ignore[attr-defined]
    for br in element.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    # Word records where it last paginated the document. Those markers are a
    # good page map when the author did not insert manual breaks.
    return any(True for _ in element.iter(qn("w:lastRenderedPageBreak")))


def _docx_blocks(document: object) -> list[tuple[str, bool]]:
    """Read paragraphs and tables in document order, with page-break flags."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks: list[tuple[str, bool]] = []
    for child in document.element.body.iterchildren():  # type: ignore[attr-defined]
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)  # type: ignore[arg-type]
            text = paragraph.text.strip()
            if text:
                blocks.append((text, _paragraph_has_page_break(paragraph)))
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)  # type: ignore[arg-type]
            for row in table.rows:
                cells: list[str] = []
                for cell in row.cells:
                    value = cell.text.strip()
                    if value and value not in cells:
                        cells.append(value)
                if cells:
                    blocks.append((" | ".join(cells), False))
    return blocks


def read_docx(data: bytes) -> Document:
    """
    Extract text from a .docx file.

    Word does not store fixed page numbers, so pages come from explicit page
    breaks (or Word's last-rendered pagination markers) when present. Without
    them every sentence has ``page=None`` and the UI says so.
    """
    if not data:
        raise DocxError("The uploaded file was empty.")
    if len(data) > MAX_DOCX_BYTES:
        raise DocxError("That .docx is larger than 15 MB. Try a smaller file or paste the text.")
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise DocxError("That is an old .doc file. Open it in Word and save it as .docx first.")
    if not data.startswith(b"PK\x03\x04"):
        raise DocxError("That file is not a .docx document. Upload a .docx file or paste the text instead.")

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            info = archive.getinfo("word/document.xml")
            if info.file_size > MAX_DOCX_XML_BYTES:
                raise DocxError(
                    "That .docx is too large to read safely. Try exporting a smaller file."
                )
            total = sum(entry.file_size for entry in archive.infolist())
            if total > MAX_DOCX_TOTAL_BYTES:
                raise DocxError(
                    "That .docx expands to too much data to read safely. Try a smaller file."
                )
    except DocxError:
        raise
    except KeyError as exc:
        raise DocxError(
            "That file is not a valid .docx document. Re-export it from Word or Google Docs."
        ) from exc
    except Exception as exc:  # noqa: BLE001 - zipfile raises several error types
        raise DocxError(
            "That file is not a valid .docx document. Re-export it from Word or Google Docs."
        ) from exc

    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover - dependency is in requirements.txt
        raise DocxError("DOCX support is not installed on the server (missing python-docx).") from exc

    try:
        document = DocxDocument(io.BytesIO(data))
        blocks = _docx_blocks(document)
    except Exception as exc:  # noqa: BLE001 - python-docx raises several different errors
        raise DocxError(
            "That .docx could not be read. It may be damaged or use an unsupported format."
        ) from exc

    if not blocks:
        raise DocxError(
            "This .docx has no selectable text. It may contain only images — paste the text instead."
        )

    page_blocks: list[tuple[int | None, str]] = []
    current: list[str] = []
    page_number = 1
    used_page_breaks = False
    for block_text, break_after in blocks:
        current.append(block_text)
        if break_after:
            used_page_breaks = True
            page_blocks.append((page_number, " ".join(current)))
            current = []
            page_number += 1
    if current:
        page_blocks.append((page_number if used_page_breaks else None, " ".join(current)))

    sentences: list[Sentence] = []
    kept_pages: list[str] = []
    total_chars = 0
    truncated = False
    for page, raw_page in page_blocks:
        cleaned = clean_text(raw_page)
        if not cleaned:
            continue
        remaining = MAX_TEXT_CHARS - total_chars
        if remaining <= 0:
            truncated = True
            break
        if len(cleaned) > remaining:
            cleaned = cleaned[:remaining]
            truncated = True
        kept_pages.append(cleaned)
        total_chars += len(cleaned)
        offset = len(sentences)
        sentences.extend(
            Sentence(text=sentence, page=page, index=offset + position)
            for position, sentence in enumerate(split_sentences(cleaned))
        )

    text = " ".join(kept_pages)
    if not sentences or len(text.strip()) < MIN_DOCX_TEXT_CHARS:
        raise DocxError(
            "This .docx has almost no selectable text. It may contain only images — paste the text instead."
        )

    warnings: list[str] = []
    if truncated:
        warnings.append(
            "Only part of the document was summarized (the limit is 120,000 characters)."
        )
    if not used_page_breaks:
        warnings.append(
            "Word documents do not store fixed page numbers, so quotes from a .docx are not tied "
            "to a page. Page numbers appear only when the file contains page breaks."
        )

    return Document(
        sentences=sentences,
        text=text,
        word_count=sum(len(sentence.text.split()) for sentence in sentences),
        kind="docx",
        page_count=max((page for page, _ in page_blocks if page is not None), default=None),
        truncated=truncated,
        warnings=warnings,
    )


def excerpt(text: str, limit: int = 700) -> str:
    cleaned = remove_citations(clean_text(text))
    if len(cleaned) <= limit:
        return cleaned
    cut = cleaned[:limit]
    tail = cut[limit // 2 :]
    if "." in tail:
        cut = cut[: limit // 2 + tail.rfind(".") + 1]
    return cut.rstrip() + " …"


def remove_citations(text: str) -> str:
    """Drop bracketed numeric citations for summary bullets (quotes keep them)."""
    return _INLINE_CITATION_RE.sub("", text).strip()
